"""Raporu Word (.docx) olarak uretir: sonuc_ve_rapor.md + Ek A (tasarim/hesaplar) + Ek B (hipotez).
Kullanim: python md2docx.py
Not: LaTeX formulleri duz metin (Unicode) olarak yazilir.
"""
import os
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

sys.stdout.reconfigure(encoding="utf-8")
ROOT = os.path.dirname(os.path.abspath(__file__))
REP = os.path.join(ROOT, "report")
OUT = os.path.join(REP, "Kovan_Rapor.docx")

SYM = [(r"\\ge(q)?", "≥"), (r"\\le(q)?", "≤"), (r"\\cdot", "·"), (r"\\times", "×"), (r"\\sum", "Σ"),
       (r"\\to", "→"), (r"\\in", "∈"), (r"\\checkmark", "✓"), (r"\\partial", "∂"), (r"\\mu", "μ"),
       (r"\\sigma", "σ"), (r"\\Delta", "Δ"), (r"\\implies", "⇒"), (r"\\;|\\,|\\!", " "),
       (r"\\\{", "{"), (r"\\\}", "}"), (r"\\\|", "‖"), (r"\\top", "ᵀ"), (r"\\approx", "≈")]


def tex(s):
    s = re.sub(r"\\underbrace\{([^{}]*)\}_\{([^{}]*)\}", r"\1", s)
    s = re.sub(r"\\hat\{(\w)\}", r"\1̂", s)
    s = re.sub(r"\\(?:text|mathrm|mathbf)\{([^{}]*)\}", r"\1", s)
    s = re.sub(r"\\frac\{([^{}]*)\}\{([^{}]*)\}", r"(\1)/(\2)", s)
    for pat, rep in SYM:
        s = re.sub(pat, rep, s)
    s = re.sub(r"\\left|\\right|\\big|\\Big", "", s)
    s = s.replace("{", "").replace("}", "").replace("\\", "")
    return re.sub(r"\s+", " ", s).strip()


TOK = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\$[^$]+\$|\*[^*\s][^*]*\*)")


def add_inline(p, text, base_bold=False):
    for part in TOK.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = p.add_run(part[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(9.5)
        elif part.startswith("$") and part.endswith("$"):
            r = p.add_run(tex(part[1:-1])); r.italic = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r = p.add_run(part[1:-1]); r.italic = True
        else:
            r = p.add_run(part)
        if base_bold:
            r.bold = True


def add_table(doc, rows):
    rows = [r for r in rows if not re.match(r"^\|?\s*:?-{2,}", r.strip())]
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    n = max(len(c) for c in cells)
    t = doc.add_table(rows=len(cells), cols=n)
    t.style = "Light Grid Accent 1"
    for i, row in enumerate(cells):
        for j in range(n):
            cell = t.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            add_inline(p, row[j] if j < len(row) else "", base_bold=(i == 0))
            for r in p.runs:
                r.font.size = Pt(9)
    doc.add_paragraph()


def render(doc, path, shift=0):
    base = os.path.dirname(path)
    lines = open(path, encoding="utf-8").read().splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i]
        s = ln.strip()
        if not s or re.fullmatch(r"-{3,}", s):
            i += 1
            continue
        if s.startswith("|"):
            blk = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                blk.append(lines[i]); i += 1
            add_table(doc, blk)
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", s)
        if m:
            lvl = min(len(m.group(1)) + shift, 4)
            doc.add_heading(tex(m.group(2)) if "$" in m.group(2) else m.group(2).replace("**", ""), level=lvl)
            i += 1
            continue
        m = re.match(r"^!\[(.*?)\]\((.*?)\)", s)
        if m:
            img = os.path.normpath(os.path.join(base, m.group(2)))
            if os.path.exists(img):
                doc.add_picture(img, width=Inches(6.0))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap = doc.add_paragraph(m.group(1)); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap.runs[0].italic = True; cap.runs[0].font.size = Pt(9)
            i += 1
            continue
        if s.startswith("$$") and s.endswith("$$") and len(s) > 4:
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(tex(s[2:-2])); r.font.name = "Cambria Math"; r.italic = True
            i += 1
            continue
        if s.startswith(">"):
            txt = re.sub(r"^>\s?", "", s)
            if re.match(r"^\[!\w+\]", txt):
                i += 1
                continue
            p = doc.add_paragraph(style="Quote")
            add_inline(p, txt)
            i += 1
            continue
        m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)", ln)
        if m:
            style = "List Number" if m.group(2)[0].isdigit() else "List Bullet"
            if len(m.group(1)) >= 2:
                style += " 2"
            p = doc.add_paragraph(style=style)
            add_inline(p, m.group(3))
            i += 1
            continue
        p = doc.add_paragraph()
        add_inline(p, s)
        i += 1


def main():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"; st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Inches(0.9)
        sec.top_margin = sec.bottom_margin = Inches(0.8)

    render(doc, os.path.join(REP, "sonuc_ve_rapor.md"))
    doc.add_page_break()
    doc.add_heading("Ek A — Tasarım: Receptive Field, Parametre ve MAC (Elle Hesap)", level=1)
    render(doc, os.path.join(REP, "tasarim_ve_hesaplamalar.md"), shift=1)
    doc.add_page_break()
    doc.add_heading("Ek B — Hipotez (Eğitimden Önce Yazılmış)", level=1)
    render(doc, os.path.join(REP, "hipotez.md"), shift=1)
    doc.save(OUT)
    print("yazildi:", OUT)


if __name__ == "__main__":
    main()
